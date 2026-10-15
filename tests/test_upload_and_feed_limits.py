import io
import os
import socket
import sys
import threading
import unittest
import zipfile
from unittest.mock import patch

import docx
import requests
from fastapi import HTTPException
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "backend"))

from routers import calendar
from services import file_extraction as files


class FakeUpload:
    def __init__(self, data, size=None):
        self.file = io.BytesIO(data)
        self.size = size


def pdf_with_pages(count):
    writer = PdfWriter()
    for _ in range(count):
        writer.add_blank_page(width=200, height=200)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def office_bytes(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class UploadLimitTests(unittest.TestCase):
    def assertStatus(self, status, call, *args):
        with self.assertRaises(HTTPException) as caught:
            call(*args)
        self.assertEqual(caught.exception.status_code, status)

    def test_declared_size_over_the_limit_is_refused_before_reading(self):
        upload = FakeUpload(b"", size=files.MAX_UPLOAD_BYTES + 1)
        self.assertStatus(413, files.read_upload, upload)

    def test_undeclared_size_stops_reading_one_byte_past_the_limit(self):
        with patch.object(files, "MAX_UPLOAD_BYTES", 10):
            self.assertStatus(413, files.read_upload, FakeUpload(b"x" * 11))
            self.assertEqual(files.read_upload(FakeUpload(b"x" * 10)), b"x" * 10)

    def test_office_file_that_unpacks_too_large_is_refused(self):
        document = docx.Document()
        document.add_paragraph("Mitosis makes two identical cells.")
        data = office_bytes(document)
        with patch.object(files, "MAX_UNZIPPED_BYTES", 1_000):
            self.assertStatus(413, files.docx_text, data)
        self.assertIn("Mitosis", files.docx_text(data))

    def test_not_a_zip_is_a_readable_error(self):
        self.assertStatus(422, files.pptx_slides, b"plain bytes, not a deck")

    def test_pdf_page_limit(self):
        with patch.object(files, "MAX_PDF_PAGES", 3):
            self.assertStatus(413, files.open_pdf, pdf_with_pages(4))
            self.assertEqual(len(files.open_pdf(pdf_with_pages(3)).pages), 3)

    def test_pptx_keeps_slide_markers_and_notes(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Cell cycle"
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Interphase comes first"
        slide.notes_slide.notes_text_frame.text = "Ask about G1"
        result = files.pptx_slides(office_bytes(deck))
        self.assertTrue(result["text"].startswith("--- Slide 1 ---"))
        self.assertEqual(result["slides"][0]["texts"], ["Cell cycle", "Interphase comes first", "Notes: Ask about G1"])

    def test_oversized_image_is_refused_before_it_reaches_the_model(self):
        with patch.object(files, "MAX_IMAGE_BYTES", 5):
            self.assertStatus(413, files.image_data_url, b"123456", "png")
        self.assertTrue(files.image_data_url(b"12", "png").startswith("data:image/png;base64,"))

    def test_binary_file_with_unknown_extension_is_refused(self):
        self.assertStatus(400, files.plain_text, bytes(range(256)) * 4, "bin")
        self.assertEqual(files.plain_text(b'{"a": 1}', "json"), '{"a": 1}')

    def test_zip_check_counts_every_entry(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            for index in range(3):
                archive.writestr(f"part{index}.xml", "0" * 600)
        with patch.object(files, "MAX_UNZIPPED_BYTES", 1_500):
            self.assertStatus(413, files._open_office_zip, out.getvalue(), "DOCX")


class CalendarFeedConnectionTests(unittest.TestCase):
    def test_connection_to_a_private_address_is_refused_after_dns(self):
        server = socket.socket()
        server.bind(("127.0.0.1", 0))
        server.listen(1)
        port = server.getsockname()[1]
        accepted = []
        threading.Thread(target=lambda: accepted.append(server.accept()[0]), daemon=True).start()
        try:
            with calendar._feed_session() as session:
                with self.assertRaises(requests.ConnectionError) as refused:
                    session.get(f"https://127.0.0.1:{port}/feed.ics", timeout=5)
            self.assertIn("non-public address", str(refused.exception))
            with self.assertRaises(HTTPException) as caught:
                with calendar._feed_session() as session:
                    calendar._follow_feed(session, f"https://127.0.0.1:{port}/feed.ics")
            self.assertEqual(caught.exception.status_code, 400)
        finally:
            for conn in accepted:
                conn.close()
            server.close()

    def test_session_ignores_environment_proxies(self):
        self.assertFalse(calendar._feed_session().trust_env)


if __name__ == "__main__":
    unittest.main()
