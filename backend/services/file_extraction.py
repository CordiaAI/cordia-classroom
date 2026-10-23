"""Read the text out of a study file a student uploads.

Every limit here exists because the input is untrusted: a small upload can
expand into gigabytes (zipped Office files) or thousands of pages (PDFs).
"""
import base64
import io
import zipfile

import docx
from fastapi import HTTPException, UploadFile
from pptx import Presentation
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
# OpenAI vision takes 20 MB per image; base64 adds a third.
MAX_IMAGE_BYTES = 15 * 1024 * 1024
MAX_UNZIPPED_BYTES = 300 * 1024 * 1024
MAX_PDF_PAGES = 600
MAX_SCANNED_PAGES = 20
MAX_TEXT_CHARS = 500_000

IMAGE_TYPES = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif",
    "webp": "image/webp", "bmp": "image/bmp", "tiff": "image/tiff", "tif": "image/tiff",
}
PLAIN_TEXT_TYPES = {"txt", "md", "csv"}
GROUP_SHAPE = 6


def _too_large(detail: str) -> HTTPException:
    return HTTPException(status_code=413, detail=detail)


def extension_of(filename: str) -> str:
    name = (filename or "").lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def read_upload(file: UploadFile) -> bytes:
    """Read the upload, refusing anything over the limit without reading all of it."""
    if file.size is not None and file.size > MAX_UPLOAD_BYTES:
        raise _too_large("File too large (max 50MB)")
    data = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise _too_large("File too large (max 50MB)")
    return data


def _open_office_zip(data: bytes, kind: str) -> None:
    """DOCX and PPTX are zip archives; check the unpacked size before parsing."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(entry.file_size for entry in archive.infolist())
    except zipfile.BadZipFile:
        raise HTTPException(status_code=422, detail=f"Could not read {kind} file.")
    if unpacked > MAX_UNZIPPED_BYTES:
        raise _too_large(f"This {kind} file is too large to read.")


def open_pdf(data: bytes) -> PdfReader:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
    except Exception:
        raise HTTPException(status_code=422, detail="Could not extract text from PDF. Use a text-based PDF.")
    if pages > MAX_PDF_PAGES:
        raise _too_large(f"This PDF has {pages} pages. Upload {MAX_PDF_PAGES} pages or fewer.")
    return reader


def pdf_text(reader: PdfReader) -> str:
    try:
        return "\n\n".join(filter(None, (page.extract_text() for page in reader.pages)))
    except Exception:
        raise HTTPException(status_code=422, detail="Could not extract text from PDF. Use a text-based PDF.")


def pdf_page_scans(reader: PdfReader) -> list:
    """The largest JPEG/PNG on each of the first pages, as data URLs, for PDFs with no text layer."""
    scans = []
    for page in reader.pages[:MAX_SCANNED_PAGES]:
        try:
            images = [image for image in page.images if extension_of(image.name) in ("jpg", "jpeg", "png")]
        except Exception:
            continue
        if not images:
            continue
        scan = max(images, key=lambda image: len(image.data))
        if len(scan.data) > MAX_IMAGE_BYTES:
            continue
        scans.append(image_data_url(scan.data, extension_of(scan.name)))
    return scans


def docx_text(data: bytes) -> str:
    _open_office_zip(data, "DOCX")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise HTTPException(status_code=422, detail="Could not read DOCX file.")
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def _slide_texts(slide) -> list:
    texts = []
    shapes = list(slide.shapes)
    while shapes:
        shape = shapes.pop(0)
        if getattr(shape, "shape_type", None) == GROUP_SHAPE:
            shapes[:0] = list(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))
        elif getattr(shape, "text", "").strip():
            texts.append(shape.text.strip())
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
    return texts


def pptx_slides(data: bytes) -> dict:
    """Slide text plus "--- Slide N ---" markers, which /generate uses to detect slideshows."""
    _open_office_zip(data, "PPTX")
    try:
        slides = [
            {"number": number, "texts": _slide_texts(slide)}
            for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1)
        ]
    except Exception:
        raise HTTPException(status_code=422, detail="Could not read PPTX file.")
    marked = "\n\n".join(
        f"--- Slide {slide['number']} ---\n" + "\n".join(slide["texts"])
        for slide in slides if slide["texts"]
    )
    return {"text": marked[:MAX_TEXT_CHARS], "slides": slides}


def plain_text(data: bytes, extension: str) -> str:
    if extension in PLAIN_TEXT_TYPES:
        return data.decode("utf-8", errors="ignore")
    # Any other extension (json, html, source code, …) is accepted only if it reads as text.
    try:
        decoded = data.decode("utf-8")
    except UnicodeDecodeError:
        decoded = data.decode("latin-1")
    printable = sum(1 for char in decoded if char.isprintable() or char in "\n\r\t")
    if not decoded or printable / len(decoded) < 0.85:
        raise HTTPException(
            status_code=400,
            detail="This file appears to be a binary file with no readable text. Please upload a PDF, DOCX, PPTX, image, or text file.",
        )
    return decoded


def text_from_file(name: str, data: bytes, mime: str = "") -> str:
    """Readable text from a downloaded file (PDF, Word, PowerPoint or text), for Library imports.
    Raises HTTPException 415 for anything else (images and scans go through the upload flow)."""
    extension = extension_of(name) or {"text/plain": "txt", "text/csv": "csv", "application/pdf": "pdf"}.get(mime, "")
    if extension == "pdf":
        return pdf_text(open_pdf(data))
    if extension == "docx":
        return docx_text(data)
    if extension == "pptx":
        return pptx_slides(data).get("text", "")
    if extension in {"txt", "md", "csv", "rtf", "html"} or mime.startswith("text/"):
        return plain_text(data, extension or "txt")
    raise HTTPException(status_code=415, detail="Cordia can import PDFs, Word, PowerPoint, Google Docs and text files.")


def image_data_url(data: bytes, extension: str) -> str:
    if len(data) > MAX_IMAGE_BYTES:
        raise _too_large("Image too large (max 15MB)")
    mime = IMAGE_TYPES.get(extension, "image/jpeg")
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"
